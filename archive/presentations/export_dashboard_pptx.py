"""
Capture screenshots of each dashboard tab and export as PowerPoint.
Usage: python export_dashboard_pptx.py
Requires: playwright, python-pptx, and a running Streamlit app on localhost:8502
"""
import time
import os
from pathlib import Path
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

APP_URL = "http://localhost:8502"
TABS = ["Dashboard", "Overview", "Analysis", "News & Sentiment", "Technical", "Fundamentals", "Backtest"]
OUT_DIR = Path("tab_screenshots")
OUT_DIR.mkdir(exist_ok=True)

def capture_tabs():
    """Use Playwright to screenshot each tab."""
    screenshots = []
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page(viewport={"width": 1400, "height": 900})
        page.goto(APP_URL, wait_until="networkidle", timeout=60000)
        # Wait for the app to fully load
        page.wait_for_timeout(8000)

        for tab_name in TABS:
            # Click the tab button
            tab_buttons = page.locator(f'button[data-baseweb="tab"]:has-text("{tab_name}")')
            if tab_buttons.count() > 0:
                tab_buttons.first.click()
                page.wait_for_timeout(4000)

            # Scroll to top
            page.evaluate("window.scrollTo(0, 0)")
            page.wait_for_timeout(500)

            # Full-page screenshot
            fname = OUT_DIR / f"{tab_name.replace(' ', '_').replace('&', 'and')}.png"
            page.screenshot(path=str(fname), full_page=True)
            screenshots.append((tab_name, str(fname)))
            print(f"  Captured: {tab_name}")

        browser.close()
    return screenshots


def build_pptx(screenshots):
    """Build a PowerPoint with one slide per tab screenshot."""
    prs = Presentation()
    # Widescreen 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    for tab_name, img_path in screenshots:
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

        # Title text box at top
        from pptx.util import Emu
        txBox = slide.shapes.add_textbox(Inches(0.3), Inches(0.15), Inches(12), Inches(0.5))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = f"{tab_name} Tab"
        p.font.size = Pt(24)
        p.font.bold = True
        from pptx.dml.color import RGBColor
        p.font.color.rgb = RGBColor(0x1A, 0x3C, 0x40)

        # Add subtitle for annotation instructions
        txBox2 = slide.shapes.add_textbox(Inches(0.3), Inches(0.6), Inches(12), Inches(0.35))
        tf2 = txBox2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = "Annotate: circle sections to resize, cross out to remove, draw arrows to reorder"
        p2.font.size = Pt(11)
        p2.font.italic = True
        from pptx.dml.color import RGBColor
        p2.font.color.rgb = RGBColor(0x5A, 0x7D, 0x82)

        # Add screenshot image — fit to slide width with margin
        from PIL import Image
        img = Image.open(img_path)
        img_w, img_h = img.size
        max_w = Inches(12.7)
        max_h = Inches(6.2)
        # Scale to fit
        scale_w = max_w / Emu(int(img_w * 914400 / 96))  # 96 dpi
        scale_h = max_h / Emu(int(img_h * 914400 / 96))
        scale = min(scale_w, scale_h, 1.0)
        final_w = Emu(int(img_w * 914400 / 96 * scale))
        final_h = Emu(int(img_h * 914400 / 96 * scale))
        left = (prs.slide_width - final_w) // 2
        top = Inches(1.0)
        slide.shapes.add_picture(img_path, left, top, final_w, final_h)

    out_path = "dashboard_layout.pptx"
    prs.save(out_path)
    print(f"\nPowerPoint saved: {os.path.abspath(out_path)}")
    return out_path


if __name__ == "__main__":
    print("Capturing tab screenshots...")
    shots = capture_tabs()
    print(f"\nBuilding PowerPoint ({len(shots)} slides)...")
    build_pptx(shots)
    print("Done!")
