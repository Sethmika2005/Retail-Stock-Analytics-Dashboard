"""Editable PowerPoint flowchart for market regime detection.

Source: models.py : detect_market_regime (lines 249-300).
Logic order (short-circuit):
  1. VIX > 25 OR VIX > 1.3 * VIX_MA20  -> High-Volatility
  2. price > SMA200 by >2% AND SMA200 slope > 0 AND SMA50 > SMA200  -> Bull
  3. price < SMA200 by >2% AND SMA200 slope < 0 AND SMA50 < SMA200  -> Bear
  4. otherwise  -> Sideways
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


BLUE = RGBColor(0x1F, 0x4E, 0x79)
BLUE_FILL = RGBColor(0xDE, 0xEB, 0xF7)
GREEN = RGBColor(0x38, 0x76, 0x1D)
GREEN_FILL = RGBColor(0xE2, 0xEF, 0xDA)
RED = RGBColor(0xC0, 0x00, 0x00)
RED_FILL = RGBColor(0xFB, 0xE5, 0xE5)
ORANGE = RGBColor(0xBF, 0x60, 0x00)
ORANGE_FILL = RGBColor(0xFD, 0xEA, 0xD0)
GREY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)


def add_shape(slide, shape_type, x, y, w, h, text, fill=WHITE, line=BLACK,
              font_size=10, bold=False, font_color=BLACK, line_width=1.0):
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    shp.line.width = Pt(line_width)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(36000); tf.margin_right = Emu(36000)
    tf.margin_top = Emu(18000); tf.margin_bottom = Emu(18000)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = "Calibri"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = font_color
    return shp


def add_label(slide, x, y, w, h, text, font_size=9, italic=True, color=GREY):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run = p.add_run(); run.text = text
    run.font.name = "Calibri"; run.font.size = Pt(font_size)
    run.font.italic = italic; run.font.color.rgb = color
    return tb


def add_title(slide, x, y, w, h, text, font_size=14, color=BLACK):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run(); run.text = text
    run.font.name = "Calibri"; run.font.size = Pt(font_size); run.font.bold = True
    run.font.color.rgb = color
    return tb


def connect(slide, src, dst, color=BLACK, width=1.25):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.ELBOW, 0, 0, 0, 0)
    conn.begin_connect(src, 2); conn.end_connect(dst, 0)
    conn.line.color.rgb = color; conn.line.width = Pt(width)
    return conn


def connect_sides(slide, src, si, dst, di, color=BLACK, width=1.25):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.ELBOW, 0, 0, 0, 0)
    conn.begin_connect(src, si); conn.end_connect(dst, di)
    conn.line.color.rgb = color; conn.line.width = Pt(width)
    return conn


def build():
    prs = Presentation()
    prs.slide_width = Inches(11)
    prs.slide_height = Inches(8.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    add_title(slide, 0.4, 0.2, 10, 0.4,
              "Market Regime Detection — Bull / Bear / Sideways / High-Volatility")

    # Inputs
    n_input = add_shape(slide, MSO_SHAPE.PARALLELOGRAM, 0.5, 1.0, 3.0, 0.6,
                        "S&P 500 OHLCV + VIX (yfinance)", fill=BLUE_FILL, line=BLUE, bold=True)
    n_compute = add_shape(slide, MSO_SHAPE.RECTANGLE, 0.5, 1.9, 3.0, 1.2,
                          "Compute:\n• SMA50, SMA200\n• SMA200 slope (20-day change)\n• price vs SMA200 (%)\n• VIX, VIX 20-day MA",
                          fill=BLUE_FILL, line=BLUE, font_size=9)
    connect(slide, n_input, n_compute, color=BLUE)

    # Decision 1: VIX
    n_d_vix = add_shape(slide, MSO_SHAPE.DIAMOND, 4.2, 1.7, 2.6, 1.5,
                        "VIX > 25\nOR\nVIX > 1.3 × VIX_MA20 ?",
                        fill=BLUE_FILL, line=BLUE, bold=True, font_size=10)
    connect_sides(slide, n_compute, 3, n_d_vix, 1, color=BLUE)

    n_high_vol = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 8.0, 2.05, 2.5, 0.8,
                           "High-Volatility",
                           fill=RED_FILL, line=RED, bold=True, font_size=13)
    connect_sides(slide, n_d_vix, 3, n_high_vol, 1, color=RED)
    add_label(slide, 6.9, 2.1, 0.9, 0.3, "Yes", color=RED, italic=False)

    # Decision 2: Bull
    n_d_bull = add_shape(slide, MSO_SHAPE.DIAMOND, 4.2, 3.8, 2.6, 1.5,
                         "price > SMA200 by > 2%\nAND SMA200 slope > 0\nAND SMA50 > SMA200 ?",
                         fill=BLUE_FILL, line=BLUE, font_size=9)
    connect(slide, n_d_vix, n_d_bull, color=BLUE)
    add_label(slide, 5.3, 3.35, 0.6, 0.3, "No")

    n_bull = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 8.0, 4.15, 2.5, 0.8,
                       "Bull",
                       fill=GREEN_FILL, line=GREEN, bold=True, font_size=13)
    connect_sides(slide, n_d_bull, 3, n_bull, 1, color=GREEN)
    add_label(slide, 6.9, 4.2, 0.9, 0.3, "Yes", color=GREEN, italic=False)

    # Decision 3: Bear
    n_d_bear = add_shape(slide, MSO_SHAPE.DIAMOND, 4.2, 5.85, 2.6, 1.5,
                         "price < SMA200 by > 2%\nAND SMA200 slope < 0\nAND SMA50 < SMA200 ?",
                         fill=BLUE_FILL, line=BLUE, font_size=9)
    connect(slide, n_d_bull, n_d_bear, color=BLUE)
    add_label(slide, 5.3, 5.40, 0.6, 0.3, "No")

    n_bear = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 8.0, 6.2, 2.5, 0.8,
                       "Bear",
                       fill=RED_FILL, line=RED, bold=True, font_size=13)
    connect_sides(slide, n_d_bear, 3, n_bear, 1, color=RED)
    add_label(slide, 6.9, 6.25, 0.9, 0.3, "Yes", color=RED, italic=False)

    # Default
    n_side = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 4.2, 7.55, 2.6, 0.7,
                       "Sideways",
                       fill=ORANGE_FILL, line=ORANGE, bold=True, font_size=13)
    connect(slide, n_d_bear, n_side, color=ORANGE)
    add_label(slide, 5.3, 7.35, 0.6, 0.25, "No (default)")

    # Legend
    add_title(slide, 0.5, 4.0, 3.2, 0.3, "Regime Implications", font_size=11, color=GREY)
    notes = [
        ("Bull", "Long-biased; rule-based BUY signals are more reliable.", GREEN, GREEN_FILL),
        ("Bear", "Short-biased; SELL/HOLD preferred. BUYs treated with caution.", RED, RED_FILL),
        ("Sideways", "Choppy; mean-reversion risk. Favour short holding periods.", ORANGE, ORANGE_FILL),
        ("High-Volatility", "Elevated risk; signals are less dependable, wider stops advised.", RED, RED_FILL),
    ]
    for i, (name, desc, stroke, fill) in enumerate(notes):
        y = 4.35 + i * 0.7
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.5, y, 1.1, 0.5, name,
                  fill=fill, line=stroke, bold=True, font_size=10)
        add_label(slide, 1.7, y + 0.05, 2.5, 0.5, desc, italic=False, font_size=8, color=BLACK)

    add_label(slide, 0.5, 8.2, 6.5, 0.25,
              "Source: models.py — detect_market_regime (lines 249–300).",
              font_size=8)

    out = Path(__file__).parent / "market_regime_flowchart.pptx"
    prs.save(out)
    print(f"Saved: {out}")


if __name__ == "__main__":
    build()
