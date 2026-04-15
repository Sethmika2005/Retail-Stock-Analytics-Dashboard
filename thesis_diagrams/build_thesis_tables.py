"""Consolidated Excel workbook for thesis Methodology chapter.

Sheet 1: Confidence Logic  (models.py : generate_recommendation_paper1, lines 449-462)
Sheet 2: Market Regime     (models.py : detect_market_regime, lines 249-300)
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Calibri", size=10)
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1F4E79")
NOTE_FONT = Font(name="Calibri", size=9, italic=True, color="595959")

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

THIN = Side(border_style="thin", color="A6A6A6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

BAND_A = PatternFill("solid", fgColor="FFFFFF")
BAND_B = PatternFill("solid", fgColor="F2F7FB")

CONF_HIGH = PatternFill("solid", fgColor="C6EFCE")
CONF_MID = PatternFill("solid", fgColor="FFEB9C")
CONF_LOW = PatternFill("solid", fgColor="FFC7CE")

FILL_HIGHVOL = PatternFill("solid", fgColor="F8CBAD")
FILL_BULL = PatternFill("solid", fgColor="C6EFCE")
FILL_BEAR = PatternFill("solid", fgColor="FFC7CE")
FILL_SIDE = PatternFill("solid", fgColor="FFEB9C")


def conf_fill(value):
    if value >= 70:
        return CONF_HIGH
    if value >= 50:
        return CONF_MID
    return CONF_LOW


def add_header_row(ws, row, headers):
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = CENTER
        c.border = BORDER


def build_confidence_sheet(ws):
    ws.title = "Confidence Logic"

    ws["A1"] = "Confidence Percentage — Decision Logic Lookup"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:F1")

    ws["A2"] = ("Confidence is derived by comparing the final rule signal (after SMA crossover, "
                "ATV confirmation, and RSI gate) against the RL agent's prediction.")
    ws["A2"].font = NOTE_FONT
    ws["A2"].alignment = LEFT
    ws.merge_cells("A2:F2")

    headers = [
        "#", "Final Rule Signal", "RL Prediction",
        "Decision Path", "Confidence (%)", "Rationale",
    ]
    add_header_row(ws, 4, headers)

    rows = [
        (1, "BUY / SELL", "Matches rule (same action)",
         "Rule fires + RL confirms", 90,
         "Rule layer produces a BUY or SELL after passing SMA crossover, ATV confirmation, and RSI gate; RL independently reaches the same action. Strongest consensus."),
        (2, "BUY / SELL", "N/A",
         "Rule fires (RL not used)", 75,
         "Rule layer produces a trade signal; no RL input is available. Confidence reflects a standalone Paper 1 decision."),
        (3, "BUY / SELL", "Opposes rule (HOLD or opposite action)",
         "Rule wins over RL dissent", 60,
         "Rule takes precedence but RL disagrees. Confidence is tempered — treat as a cautious entry."),
        (4, "HOLD", "HOLD (agrees)",
         "Consensus HOLD", 70,
         "Both layers independently arrive at HOLD. High confidence in the no-trade decision."),
        (5, "HOLD", "BUY / SELL (opposes rule)",
         "RL dissents from HOLD", 55,
         "Rule says HOLD but RL sees an opportunity. Where rule had no opinion (no crossover), RL overrides and the recommendation becomes the RL signal."),
        (6, "HOLD", "N/A",
         "Default HOLD", 40,
         "Rule produces HOLD and no RL input is available. Lowest actionable confidence — essentially a no-signal state."),
    ]

    for i, row in enumerate(rows):
        r = 5 + i
        band = BAND_A if i % 2 == 0 else BAND_B
        for col, value in enumerate(row, start=1):
            c = ws.cell(row=r, column=col, value=value)
            c.font = BODY_FONT
            c.alignment = CENTER if col != 6 else LEFT
            c.border = BORDER
            c.fill = band
        ws.cell(row=r, column=5).fill = conf_fill(row[4])
        ws.cell(row=r, column=5).font = Font(name="Calibri", size=11, bold=True)

    widths = [4, 18, 32, 28, 14, 60]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(5, 5 + len(rows)):
        ws.row_dimensions[r].height = 48
    ws.row_dimensions[4].height = 32

    legend_row = 5 + len(rows) + 2
    ws.cell(row=legend_row, column=1, value="Confidence bands:").font = Font(
        name="Calibri", size=10, bold=True
    )
    ws.merge_cells(start_row=legend_row, start_column=1, end_row=legend_row, end_column=2)
    for j, (label, fill) in enumerate([
        ("High (≥ 70)", CONF_HIGH),
        ("Moderate (50–69)", CONF_MID),
        ("Low (< 50)", CONF_LOW),
    ]):
        c = ws.cell(row=legend_row + 1 + j, column=1, value=label)
        c.fill = fill; c.alignment = CENTER; c.border = BORDER; c.font = BODY_FONT

    note_row = legend_row + 5
    ws.cell(row=note_row, column=1,
            value="Source: models.py — generate_recommendation_paper1 (lines 449–462).").font = NOTE_FONT
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=6)


def build_regime_sheet(ws):
    ws.title = "Market Regime"

    ws["A1"] = "Market Regime Classification — Rule Lookup"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:H1")

    ws["A2"] = ("Regimes are evaluated top-down in priority order; the first rule whose "
                "conditions are all satisfied wins. Inputs are computed from S&P 500 daily "
                "OHLCV and the CBOE VIX index.")
    ws["A2"].font = NOTE_FONT
    ws["A2"].alignment = LEFT
    ws.merge_cells("A2:H2")

    headers = [
        "Priority", "Regime", "VIX Condition", "Price vs SMA200",
        "SMA200 Slope (20-day)", "SMA50 vs SMA200", "All Conditions?", "Implication for Trading",
    ]
    add_header_row(ws, 4, headers)

    rows = [
        (1, "High-Volatility",
         "VIX > 25   OR   VIX > 1.3 × VIX_MA20",
         "—", "—", "—",
         "Either VIX clause true",
         "Elevated risk; signals less reliable. Widen stops, reduce position size, prefer HOLD.",
         FILL_HIGHVOL),
        (2, "Bull",
         "—  (VIX not elevated)",
         "Price > SMA200 by > 2%",
         "Slope > 0  (rising)",
         "SMA50 > SMA200  (golden alignment)",
         "All three true",
         "Long-biased market. Rule-based BUY signals are more reliable; trend-following favoured.",
         FILL_BULL),
        (3, "Bear",
         "—  (VIX not elevated)",
         "Price < SMA200 by > 2%",
         "Slope < 0  (falling)",
         "SMA50 < SMA200  (death alignment)",
         "All three true",
         "Short-biased market. SELL / HOLD preferred; BUY signals treated with scepticism.",
         FILL_BEAR),
        (4, "Sideways",
         "—", "—", "—", "—",
         "Default — nothing above matched",
         "Choppy / range-bound. Mean-reversion risk; favour short holding periods and tight risk limits.",
         FILL_SIDE),
    ]

    for i, row in enumerate(rows):
        r = 5 + i
        priority, regime, vix, price, slope, sma, logic, impl, regime_fill = row
        values = [priority, regime, vix, price, slope, sma, logic, impl]
        band = BAND_A if i % 2 == 0 else BAND_B
        for col, value in enumerate(values, start=1):
            c = ws.cell(row=r, column=col, value=value)
            c.font = BODY_FONT
            c.alignment = CENTER if col != 8 else LEFT
            c.border = BORDER
            c.fill = band
        ws.cell(row=r, column=2).fill = regime_fill
        ws.cell(row=r, column=2).font = Font(name="Calibri", size=11, bold=True)

    widths = [9, 16, 34, 22, 22, 26, 24, 50]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(5, 5 + len(rows)):
        ws.row_dimensions[r].height = 58
    ws.row_dimensions[4].height = 36

    legend_row = 5 + len(rows) + 2
    ws.cell(row=legend_row, column=1, value="Derived quantities (definitions):").font = Font(
        name="Calibri", size=10, bold=True
    )
    ws.merge_cells(start_row=legend_row, start_column=1, end_row=legend_row, end_column=3)

    for j, (term, desc) in enumerate([
        ("SMA200", "200-day simple moving average of S&P 500 close."),
        ("SMA50", "50-day simple moving average of S&P 500 close."),
        ("SMA200 slope", "Percent change of SMA200 over the last 20 trading days."),
        ("Price vs SMA200", "(current price − SMA200) / SMA200 × 100, in percent."),
        ("VIX_MA20", "20-day simple moving average of the CBOE VIX close."),
    ]):
        r = legend_row + 1 + j
        tcell = ws.cell(row=r, column=1, value=term)
        tcell.font = Font(name="Calibri", size=10, bold=True)
        tcell.alignment = LEFT; tcell.border = BORDER
        dcell = ws.cell(row=r, column=2, value=desc)
        dcell.font = BODY_FONT; dcell.alignment = LEFT; dcell.border = BORDER
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)

    note_row = legend_row + 7
    ws.cell(row=note_row, column=1,
            value="Source: models.py — detect_market_regime (lines 249–300).").font = NOTE_FONT
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=8)


def build():
    wb = Workbook()
    build_confidence_sheet(wb.active)
    build_regime_sheet(wb.create_sheet())

    out = Path(__file__).parent / "thesis_tables.xlsx"
    wb.save(out)
    print(f"Saved: {out}")


if __name__ == "__main__":
    build()
